"""Phase 7: presentations, snapshot, sharing, insights and the daily-log summary -- permissions and behaviour."""

from __future__ import annotations

from decimal import Decimal
from typing import Any

import httpx
import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker
from tt_schema.base import new_id
from tt_schema.enums import HabitStatus, Polarity, PolicyCategory, TargetType, Trend
from tt_server.db import models as m
from tt_server.integrations.llm import LLMGatewayClient

from tests.api.conftest import Api
from tests.api.test_acl_llm import answer, fake_gateway

ROUTERS_COVERED = ("presentations",)


async def mk(api: Api, as_: str = "@mahmoud", **body: Any) -> dict[str, Any]:
    r = await api.call("POST", "/presentations", as_=as_, json={"title": "Deck", **body})
    assert r.status_code == 201, r.text
    return r.json()  # type: ignore[no-any-return]


async def test_a_presentation_is_private_until_shared_and_sharing_needs_view_share(api: Api) -> None:
    await api.role("@sara", "contributor")
    p = await mk(api, as_="@sara", slides=[{"kind": "kpis"}])
    assert (await api.call("GET", f"/presentations/{p['id']}", as_="@ali")).status_code == 404
    assert (await api.call("GET", f"/presentations/{p['id']}", as_="@mahmoud")).status_code == 200  # admin
    shared = await api.call("PATCH", f"/presentations/{p['id']}", as_="@sara", json={"is_shared": True})
    assert shared.status_code == 403 and shared.json()["error"]["details"]["permission"] == "view.share"
    await api.role("@sara", "manager")
    assert (
        await api.call("PATCH", f"/presentations/{p['id']}", as_="@sara", json={"is_shared": True})
    ).status_code == 200
    assert (await api.call("GET", f"/presentations/{p['id']}", as_="@ali")).status_code == 200
    # only the owner (or an admin) edits
    assert (
        await api.call("PATCH", f"/presentations/{p['id']}", as_="@ali", json={"title": "mine now"})
    ).status_code == 403


async def test_slides_are_validated_per_kind_and_templates_expand(api: Api) -> None:
    bad = await api.call("POST", "/presentations", json={"title": "x", "slides": [{"kind": "view"}]})
    assert bad.status_code == 400 and bad.json()["error"]["pointer"] == "/slides/0/config/saved_view_id"
    assert (
        await api.call("POST", "/presentations", json={"title": "x", "slides": [{"kind": "nope"}]})
    ).status_code == 400
    ex = await api.call(
        "POST",
        "/presentations",
        json={"title": "x", "slides": [{"kind": "explorer", "config": {"dimA": "bogus"}}]},
    )
    assert ex.status_code == 400
    r = await api.call("POST", "/presentations:from-template", json={"template": "portfolio_review"})
    assert r.status_code == 201
    kinds = [s["kind"] for s in r.json()["slides"]]
    assert kinds == [
        "kpis",
        "projects",
        "pipeline",
        "explorer",
        "radar",
        "gantt",
        "now",
        "people",
        "violations",
    ]
    assert [s["position"] for s in r.json()["slides"]] == list(range(9))
    assert (
        await api.call("POST", "/presentations:from-template", json={"template": "nope"})
    ).status_code == 404
    names = {t["key"] for t in (await api.call("GET", "/presentations/templates")).json()["data"]}
    assert names == {"portfolio_review", "sprint_review", "risk_review", "team_1on1", "exec_summary"}


async def test_slide_add_patch_reorder_duplicate_delete(api: Api) -> None:
    p = await mk(api, slides=[{"kind": "kpis"}, {"kind": "now"}])
    a, b = (s["id"] for s in p["slides"])
    added = (
        await api.call(
            "POST",
            f"/presentations/{p['id']}/slides",
            json={"slide": {"kind": "markdown", "config": {"text": "hi"}}, "position": 0},
        )
    ).json()
    assert [s["kind"] for s in added["slides"]] == ["markdown", "kpis", "now"]
    md = added["slides"][0]["id"]
    r = await api.call("PUT", f"/presentations/{p['id']}/slides:reorder", json={"ids": [b, md, a]})
    assert [s["id"] for s in r.json()["slides"]] == [b, md, a]
    assert (
        await api.call("PUT", f"/presentations/{p['id']}/slides:reorder", json={"ids": [b]})
    ).status_code == 400
    r = await api.call(
        "PATCH", f"/presentations/{p['id']}/slides/{md}", json={"notes": "say this", "config": {"text": "yo"}}
    )
    assert r.json()["slides"][1]["notes"] == "say this" and r.json()["slides"][1]["config"]["text"] == "yo"
    dup = (await api.call("POST", f"/presentations/{p['id']}:duplicate")).json()
    assert dup["title"].endswith("(copy)") and {s["id"] for s in dup["slides"]}.isdisjoint({a, b, md})
    gone = await api.call("DELETE", f"/presentations/{p['id']}/slides/{md}")
    assert [s["kind"] for s in gone.json()["slides"]] == ["now", "kpis"]
    assert (await api.call("DELETE", f"/presentations/{p['id']}")).status_code == 204
    assert (await api.call("GET", f"/presentations/{p['id']}")).status_code == 404


async def test_snapshot_is_scoped_acl_filtered_and_matches_its_schema(api: Api) -> None:
    from jsonschema import validate

    for t, proj in (("a", "WEB"), ("b", "WEB"), ("c", "DATA")):
        await api.call("POST", "/items", json={"title": f"item {t}", "project": proj})
    snap = (await api.call("GET", "/snapshot")).json()
    schema = (await api.call("GET", "/meta/schema/snapshot")).json()
    validate(snap, schema)
    assert len(snap["items"]) == 3 and snap["workspace"] == "acme" and snap["meta"]["viewer"] == "@mahmoud"
    web = (await api.call("GET", "/snapshot", params={"scope": "project:WEB"})).json()
    assert {i["project"] for i in web["items"]} == {"WEB"} and web["meta"][
        "scope_label"
    ] == "Project: Website Refresh"
    ali = (await api.call("GET", "/snapshot", as_="@ali")).json()
    assert ali["items"] == [] and ali["projects"] == []  # not a member of anything
    assert (await api.call("GET", "/snapshot", params={"scope": "project:NOPE"})).status_code == 404
    assert (await api.call("GET", "/snapshot", params={"scope": "galaxy:1"})).status_code == 400
    assert (await api.call("GET", "/snapshot", params={"days": 0})).status_code == 400
    # cached per viewer, `fresh=1` bypasses
    again = (await api.call("GET", "/snapshot")).json()
    assert again["items"] == snap["items"]
    await api.call("POST", "/items", json={"title": "item d", "project": "WEB"})
    assert len((await api.call("GET", "/snapshot")).json()["items"]) == 3
    assert len((await api.call("GET", "/snapshot", params={"fresh": 1})).json()["items"]) == 4


async def test_inbox_items_are_not_in_a_project_managers_snapshot(api: Api) -> None:
    people = {p["handle"]: p["person_id"] for p in (await api.call("GET", "/people")).json()["data"]}
    await api.role("@ali", "manager")
    await api.send(as_="@sara", action="create", item={"title": "sara inbox note"}, project="nowhere")
    snap = (await api.call("GET", "/snapshot", as_="@ali")).json()
    assert "sara inbox note" not in {i["title"] for i in snap["items"]}
    line = {"person_id": people["@sara"], "manager_id": people["@mahmoud"], "kind": "line"}
    assert (await api.call("POST", "/reporting-lines", json=line)).status_code == 201
    # even the line manager gets inbox items only through the explicit query flag, never through the snapshot
    assert "sara inbox note" not in {i["title"] for i in (await api.call("GET", "/snapshot")).json()["items"]}


def _habit(ws_id: str, person_id: str, key: str, category: PolicyCategory) -> m.Habit:
    return m.Habit(
        id=new_id(), workspace_id=ws_id, subject_type=TargetType.PERSON, subject_id=person_id, key=key,
        polarity=Polarity.NEGATIVE, category=category, trend=Trend.STABLE, status=HabitStatus.OBSERVED,
        occurrences=3, rate=Decimal("0.5"), score=Decimal("1"),
        window="30d",
    )  # fmt: skip


async def test_wellbeing_habits_are_redacted_for_a_project_manager_but_not_a_line_manager(
    api: Api, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    people = {p["handle"]: p["person_id"] for p in (await api.call("GET", "/people")).json()["data"]}
    line = {"person_id": people["@sara"], "manager_id": people["@mahmoud"], "kind": "line"}
    assert (await api.call("POST", "/reporting-lines", json=line)).status_code == 201
    await api.call("POST", "/projects/WEB/members", json={"handle": "@sara"})
    await api.call("POST", "/projects/WEB/members", json={"handle": "@ali", "role": "manager"})
    await api.role("@ali", "manager", scope_type="project", scope_id="WEB")
    async with session_factory() as s:
        ws = await s.scalar(select(m.Workspace).where(m.Workspace.slug == "acme"))
        assert ws
        s.add(_habit(ws.id, people["@sara"], "no_estimates", PolicyCategory.PLANNING))
        s.add(_habit(ws.id, people["@sara"], "late_nights", PolicyCategory.WELLBEING))
        await s.commit()
    pm = {
        h["key"]
        for h in (await api.call("GET", "/snapshot", as_="@ali", params={"fresh": 1})).json()["habits"]
    }
    lm = {h["key"] for h in (await api.call("GET", "/snapshot", params={"fresh": 1})).json()["habits"]}
    assert pm == {"no_estimates"} and lm == {"no_estimates", "late_nights"}


async def test_data_endpoint_marks_a_restricted_habits_slide(
    api: Api, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    people = {p["handle"]: p["person_id"] for p in (await api.call("GET", "/people")).json()["data"]}
    async with session_factory() as s:
        ws = await s.scalar(select(m.Workspace).where(m.Workspace.slug == "acme"))
        assert ws
        s.add(_habit(ws.id, people["@sara"], "late_nights", PolicyCategory.WELLBEING))  # only sara may see it
        await s.commit()
    p = await mk(api, is_shared=True, slides=[{"kind": "habits"}, {"kind": "kpis"}])
    habits_id = p["slides"][0]["id"]
    ali = (await api.call("GET", f"/presentations/{p['id']}/data", as_="@ali")).json()
    assert ali["slide_data"][habits_id] == {"restricted": True} and ali["snapshot"]["habits"] == []
    sara = (await api.call("GET", f"/presentations/{p['id']}/data", as_="@sara")).json()
    assert sara["slide_data"][habits_id] == {"restricted": False}
    assert [h["key"] for h in sara["snapshot"]["habits"]] == ["late_nights"]


async def test_view_and_table_slides_carry_the_viewers_own_query_results(api: Api) -> None:
    await api.call("POST", "/items", json={"title": "wip secret", "project": "WEB"})
    view = (
        await api.call("POST", "/views", json={"name": "all web", "filter": {"project": {"in": ["WEB"]}}})
    ).json()
    p = await mk(
        api,
        is_shared=True,
        slides=[
            {"kind": "view", "config": {"saved_view_id": view["id"]}},
            {"kind": "table", "config": {"filter": {"text": "secret"}}},
            {"kind": "chart", "config": {"rows": "project"}},
        ],
    )
    sd = (await api.call("GET", f"/presentations/{p['id']}/data")).json()["slide_data"]
    v, t, c = (sd[s["id"]] for s in p["slides"])
    assert v["layout"] == "list" and [i["title"] for i in v["items"]] == ["wip secret"]
    assert t["page"]["total"] == 1 and c["totals"]["all"] == 1
    other = (await api.call("GET", f"/presentations/{p['id']}/data", as_="@ali")).json()["slide_data"]
    assert other[p["slides"][0]["id"]] == {
        "restricted": True
    }  # the owner's private view is not visible to @ali
    assert other[p["slides"][1]["id"]]["items"] == [] and other[p["slides"][2]["id"]]["totals"]["all"] == 0


async def test_external_share_token_needs_the_workspace_flag_and_expires(
    api: Api, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    p = await mk(api, slides=[{"kind": "kpis"}])
    off = await api.call("POST", f"/presentations/{p['id']}/share", json={})
    assert off.status_code == 409 and off.json()["error"]["code"] == "not_enabled"
    async with session_factory() as s:
        ws = await s.scalar(select(m.Workspace).where(m.Workspace.slug == "acme"))
        assert ws
        ws.settings = {**ws.settings, "external_share": True}
        await s.commit()
    on = (await api.call("POST", f"/presentations/{p['id']}/share", json={"expires_in_days": 7})).json()
    token = on["token"]
    assert token.startswith("pt_") and on["url"] == f"/p/t/{token}"
    anon = await api.c.get(f"/api/v1/public/p/{token}/data")  # no credentials at all
    assert anon.status_code == 200 and "owner_id" not in anon.json()["presentation"]
    assert (await api.c.get("/api/v1/public/p/pt_wrong/data")).status_code == 404
    revoked = await api.call("POST", f"/presentations/{p['id']}/share", json={"enabled": False})
    assert revoked.json()["token"] is None
    assert (await api.c.get(f"/api/v1/public/p/{token}/data")).status_code == 404
    again = (await api.call("POST", f"/presentations/{p['id']}/share", json={"expires_in_days": 1})).json()[
        "token"
    ]
    async with session_factory() as s:
        row = await s.scalar(select(m.Presentation).where(m.Presentation.share_token == again))
        assert row
        from datetime import timedelta

        from tt_schema.base import utcnow

        row.share_expires_at = utcnow() - timedelta(minutes=1)
        await s.commit()
    assert (await api.c.get(f"/api/v1/public/p/{again}/data")).status_code == 404


async def test_exports_html_csv_pptx_are_stamped_and_pdf_png_need_the_bundle(
    api: Api, tmp_path: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    import io
    import zipfile

    from pptx import Presentation as Pptx
    from tt_server.api.routers import presentations as router

    await api.call("POST", "/items", json={"title": "slide me", "project": "WEB", "priority": "high"})
    p = await mk(
        api,
        slides=[
            {"kind": "title", "title": "Q4"},
            {"kind": "kpis", "notes": "talk track"},
            {"kind": "explorer", "config": {"dimA": "project", "dimB": None, "measure": "count"}},
            {"kind": "table", "config": {"filter": {}}},
            {"kind": "markdown", "config": {"text": "- a\n- b"}},
        ],
    )
    pptx = await api.call("POST", f"/presentations/{p['id']}/render", params={"format": "pptx"})
    assert pptx.status_code == 200
    prs = Pptx(io.BytesIO(pptx.content))
    assert len(prs.slides) == 5
    assert any(
        sh.has_chart for sh in prs.slides[2].shapes
    )  # the one-dimension explorer is a native bar chart
    assert prs.slides[1].notes_slide.notes_text_frame.text == "talk track"
    footer = " ".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
    assert "Exported by @mahmoud" in footer and footer.count("· 2/5") == 1
    csvz = zipfile.ZipFile(
        io.BytesIO(
            (await api.call("POST", f"/presentations/{p['id']}/render", params={"format": "csv"})).content
        )
    )
    names = csvz.namelist()
    assert "stamp.txt" in names and any(n.endswith("table.csv") for n in names)
    from tt_server.acl.export import is_stamped

    assert is_stamped(csvz.read(next(n for n in names if n.endswith("table.csv"))).decode())
    # the HTML export needs the built bundle; without it the API says so instead of failing obscurely
    monkeypatch.setattr(
        router, "BUNDLE", tmp_path / "missing.html"
    )  # restored after the test (it used to leak a stub)
    miss = await api.call("POST", f"/presentations/{p['id']}/render", params={"format": "html"})
    assert miss.status_code == 503 and miss.json()["error"]["code"] == "bundle_missing"
    (tmp_path / "b.html").write_text("<html><head></head><body></body></html>")
    monkeypatch.setattr(router, "BUNDLE", tmp_path / "b.html")
    html = (await api.call("POST", f"/presentations/{p['id']}/render", params={"format": "html"})).text
    assert "window.__DECK__=" in html and "</script><script" not in html and "Exported by @mahmoud" in html
    assert (
        await api.call("POST", f"/presentations/{p['id']}/render", params={"format": "gif"})
    ).status_code == 400


# ---------------------------------------------------------------- insights
async def _gw(api: Api, output: Any) -> list[dict[str, Any]]:
    seen: list[dict[str, Any]] = []

    def handler(req: httpx.Request) -> httpx.Response:
        import json

        seen.append(json.loads(req.content))
        return answer(output)

    api.c._transport.app.state.llm = fake_gateway(handler)  # type: ignore[attr-defined]
    return seen


async def test_insight_without_a_gateway_uses_the_deterministic_fallback(api: Api) -> None:
    api.c._transport.app.state.llm = None  # type: ignore[attr-defined]
    await api.call("POST", "/items", json={"title": "one", "project": "WEB"})
    r = await api.call("POST", "/insights:generate", json={"scope": "WEB"})
    assert r.status_code == 201
    ins = r.json()["data"][0]
    assert (
        ins["source"] == "fallback"
        and ins["body"].startswith("0/1 items done (0 %).")
        and ins["status"] == "new"
    )
    snap = (await api.call("GET", "/snapshot", params={"fresh": 1})).json()
    assert next(p for p in snap["projects"] if p["key"] == "WEB")["summary"] == ins["body"]


async def test_grounded_gateway_summary_is_stored_and_ungrounded_one_is_rejected(api: Api) -> None:
    item = (await api.call("POST", "/items", json={"title": "one", "project": "WEB"})).json()
    seen = await _gw(api, {"body": "WEB is quiet.", "supporting_ids": ["WEB"]})
    ok = (await api.call("POST", "/insights:generate", json={"scope": "WEB"})).json()["data"][0]
    assert ok["source"] == "gateway" and ok["body"] == "WEB is quiet." and ok["supporting_ids"] == ["WEB"]
    assert seen[0]["facts"]["project"]["key"] == "WEB" and seen[0]["workspace"] == "acme"
    await _gw(api, {"body": "Invented claim about OTHER-9.", "supporting_ids": ["OTHER-9", item["key"]]})
    bad = (await api.call("POST", "/insights:generate", json={"scope": "WEB"})).json()["data"][0]
    assert bad["source"] == "fallback" and "Invented" not in bad["body"]


async def test_insight_status_workflow_and_visibility(api: Api) -> None:
    api.c._transport.app.state.llm = None  # type: ignore[attr-defined]
    ins = (await api.call("POST", "/insights:generate", json={"scope": "WEB"})).json()["data"][0]
    assert (await api.call("GET", "/insights", as_="@ali")).json()["data"] == []
    assert (
        await api.call("PATCH", f"/insights/{ins['id']}", as_="@ali", json={"status": "reviewed"})
    ).status_code == 404
    done = await api.call("PATCH", f"/insights/{ins['id']}", json={"status": "dismissed"})
    assert done.json()["status"] == "dismissed"
    assert (await api.call("GET", "/insights", params={"status": "new"})).json()["data"] == []
    # a dismissed insight no longer feeds the project card: the card falls back to the deterministic text
    snap = (await api.call("GET", "/snapshot", params={"fresh": 1})).json()
    card = next(p for p in snap["projects"] if p["key"] == "WEB")["summary"]
    assert card.startswith("0/0 items done")
    types = {e["event_type"] for e in (await api.call("GET", "/events", params={"limit": 50})).json()["data"]}
    assert {"insight.generated", "insight.status_changed"} <= types


async def test_daily_log_summary_goes_through_the_gateway_and_keeps_grounding(
    api: Api, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    from datetime import date

    from tt_server.domain.capture import dailylog

    res = await api.send(action="start", item={"title": "write docs"}, project="WEB", minutes=45)
    key = res["item"]["key"]
    async with session_factory() as s:
        ws = await s.scalar(select(m.Workspace).where(m.Workspace.slug == "acme"))
        assert ws
    await dailylog.build_for_workspace(session_factory, ws.id, date(2026, 10, 23))
    log = (await api.call("GET", "/daily-logs", params={"day": "2026-10-23"})).json()["data"][0]
    lid, raw = log["id"], log["summary"]
    api.c._transport.app.state.llm = None  # type: ignore[attr-defined]
    assert (await api.call("POST", f"/daily-logs/{lid}/summarise")).status_code == 409
    await _gw(api, {"body": "Documented the thing.", "supporting_ids": [key]})
    ok = (await api.call("POST", f"/daily-logs/{lid}/summarise")).json()
    assert ok["source"] == "gateway" and ok["summary"] == "Documented the thing."
    await _gw(api, {"body": "Did secret work.", "supporting_ids": ["ZZZ-1"]})
    kept = (await api.call("POST", f"/daily-logs/{lid}/summarise")).json()
    assert kept["source"] == "deterministic" and kept["summary"] == "Documented the thing." != raw
    assert (await api.call("POST", f"/daily-logs/{lid}/summarise", as_="@sara")).status_code == 404
    assert (await api.call("POST", f"/daily-logs/{lid}/confirm")).status_code == 200
    assert (
        await api.call("POST", f"/daily-logs/{lid}/summarise")
    ).status_code == 409  # confirmed logs are final


def test_insight_helpers_are_pure() -> None:
    from datetime import date

    from tt_server.domain import insights

    proj = {"key": "P", "name": "P"}
    items = [
        {"key": "P-1", "project": "P", "state": "done", "stage": "done", "stage_age_days": 1, "estimate_h": 1, "evidence_count": 1},
        {"key": "P-2", "project": "P", "state": "active", "stage": "test", "stage_age_days": 9, "estimate_h": None, "evidence_count": 0},
    ]  # fmt: skip
    facts = insights.project_facts(
        proj,
        items,
        [{"name": "Beta", "project": "P", "due": "2026-11-01"}],
        [{"key": "P-R1", "project": "P", "p": 4, "i": 5, "status": "open"}],
        {"test": 5},
        date(2026, 10, 23),
    )
    assert facts["counts"] == {"total": 2, "done": 1, "stuck": 1} and facts["stuck_stage"] == "test"
    text = insights.fallback_summary(facts)
    assert (
        text
        == "1/2 items done (50 %). 1 items over expected stage time, mostly in test. Next milestone Beta on 2026-11-01. Top risk P-R1 (20)."
    )
    assert insights.grounded(["P", "P-2", "P-R1"], facts) and not insights.grounded(["P-9"], facts)


_ = (LLMGatewayClient,)


async def test_snapshot_cache_refresh_job_warms_presentations_and_refreshes_stale_rows(
    api: Api, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    from datetime import timedelta

    from tt_schema.base import utcnow
    from tt_server.domain import deckdata

    await api.call("POST", "/items", json={"title": "cache me", "project": "WEB"})
    await mk(api, scope="project:WEB", slides=[{"kind": "kpis"}])  # saved, never opened: the job warms it
    await api.call("GET", "/snapshot", as_="@sara")
    async with session_factory() as s:
        ws = await s.scalar(select(m.Workspace).where(m.Workspace.slug == "acme"))
        assert ws
        rows = (await s.scalars(select(m.SnapshotCache))).all()
        assert len(rows) == 1
        rows[0].as_of = utcnow() - timedelta(hours=1)
        await s.commit()
    assert await deckdata.refresh_cache(session_factory, ws.id) == 2  # sara's stale row + the owner's warm-up
    async with session_factory() as s:
        rows = (await s.scalars(select(m.SnapshotCache))).all()
        assert len(rows) == 2 and all(utcnow() - r.as_of < timedelta(minutes=1) for r in rows)
    assert await deckdata.refresh_cache(session_factory, ws.id) == 0  # nothing is stale any more


async def test_nightly_insights_store_only_grounded_gateway_summaries(
    api: Api, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    from tt_server.domain import insights

    await api.call("POST", "/items", json={"title": "one", "project": "WEB"})
    async with session_factory() as s:
        ws_id = await s.scalar(select(m.Workspace.id).where(m.Workspace.slug == "acme"))
    assert ws_id
    await _gw(api, {"body": "Quiet.", "supporting_ids": ["WEB"]})
    gw = api.c._transport.app.state.llm  # type: ignore[attr-defined]
    assert (
        await insights.nightly(session_factory, gw, ws_id) == 1
    )  # only WEB: "WEB" is not a fact of DATA, so DATA falls back
    await _gw(api, {"body": "Made up.", "supporting_ids": ["NOPE-1"]})
    gw = api.c._transport.app.state.llm  # type: ignore[attr-defined]
    assert await insights.nightly(session_factory, gw, ws_id) == 0  # ungrounded: fallback, nothing stored
    listed = (await api.call("GET", "/insights")).json()["data"]
    assert [(i["project"], i["body"]) for i in listed] == [("WEB", "Quiet.")]


def test_jobs_are_registered() -> None:
    from tt_server.jobs import llm_tasks, tasks

    assert (
        tasks.refresh_snapshots.name == "snapshot.refresh"
        and llm_tasks.nightly_insights.name == "insights.nightly"
    )
