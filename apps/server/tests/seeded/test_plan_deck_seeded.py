"""Package F: a deck over the plan-only project ERP renders every slide, exports, and shows progress once work is reported."""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path

import pytest
from pptx import Presentation as Pptx
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from tests.api.conftest import Api

ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT))
from tools.seed.demo import persist_plan_only  # noqa: E402

BUNDLE = ROOT / "apps/presentation/dist/presentation.html"


@pytest.fixture
async def erp(seed_sf: async_sessionmaker[AsyncSession], sapi: Api) -> Api:
    await persist_plan_only(seed_sf)
    return sapi


async def deck(api: Api, template: str = "project_plan_review") -> dict[str, object]:
    r = await api.call(
        "POST", "/presentations:from-template", json={"template": template, "scope": "project:ERP"}
    )
    assert r.status_code == 201, r.text
    return r.json()  # type: ignore[no-any-return]


async def test_the_snapshot_carries_the_plan(erp: Api) -> None:
    s = (await erp.call("GET", "/snapshot", params={"scope": "project:ERP", "fresh": 1})).json()
    assert len(s["phases"]) == 3 and len(s["deliverables"]) == 9 and len(s["milestones"]) == 4
    assert {m["kind"] for m in s["milestones"]} == {"gate", "release"}
    p = s["projects"][0]
    assert (
        p["key"] == "ERP"
        and p["owner"] == "@mahmoud"
        and p["baseline_name"] == "approved"
        and p["health"] in {"green", "amber", "red"}
    )
    assert s["items"] == []  # plan containers are not "work"
    d = next(x for x in s["deliverables"] if x["title"] == "Infrastructure and environments")
    assert d["depends_on"] and d["phase"] and d["baseline_variance_days"] == 0 and d["progress_pct"] is None
    from jsonschema import validate

    validate(s, __import__("json").loads((ROOT / "docs/api/snapshot-schema.json").read_text()))


async def test_the_plan_review_template_has_the_specified_slides(erp: Api) -> None:
    p = await deck(erp)
    assert [s["kind"] for s in p["slides"]][:5] == [
        "title",
        "project_overview",
        "roadmap",
        "milestones",
        "deliverables",
    ]  # type: ignore[index]
    prog = await deck(erp, "programme_roadmap")
    assert "roadmap" in [s["kind"] for s in prog["slides"]]  # type: ignore[index]


@pytest.mark.skipif(not BUNDLE.exists(), reason="presentation bundle not built")
async def test_every_slide_exports_to_pdf_and_pptx(erp: Api) -> None:
    p = await deck(erp)
    n = len(p["slides"])  # type: ignore[arg-type]
    pdf = await erp.call("POST", f"/presentations/{p['id']}/render", params={"format": "pdf"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert len(re.findall(rb"/Type\s*/Page[^s]", pdf.content)) >= n
    pptx = await erp.call("POST", f"/presentations/{p['id']}/render", params={"format": "pptx"})
    assert pptx.status_code == 200
    prs = Pptx(io.BytesIO(pptx.content))
    text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert (
        len(prs.slides) >= n
        and "ERP Rollout" in text
        and "No work reported" not in text.replace("No work reported yet", "")
    )
    assert "General ledger configuration" in " ".join(
        c.text
        for sl in prs.slides
        for sh in sl.shapes
        if sh.has_table
        for r in sh.table.rows
        for c in r.cells
    )


async def test_progress_appears_after_work_is_logged_without_editing_the_deck(erp: Api) -> None:
    p = await deck(erp)
    before = (await erp.call("GET", f"/presentations/{p['id']}/data", params={"fresh": 1})).json()["snapshot"]
    gl = next(d for d in before["deliverables"] if d["title"] == "General ledger configuration")
    assert gl["progress_pct"] is None and gl["linked_tasks"] == 0
    a = await erp.send(
        action="create", item={"title": "configure the general ledger chart of accounts"}, project="ERP"
    )
    b = await erp.send(
        action="create", item={"title": "general ledger configuration: set up currencies"}, project="ERP"
    )
    assert a["resolution"]["plan_link"].startswith("linked:") and b["resolution"]["plan_link"].startswith(
        "linked:"
    )
    await erp.send(action="done", item={"key": a["item"]["key"]}, minutes=90)
    after = (await erp.call("GET", f"/presentations/{p['id']}/data", params={"fresh": 1})).json()["snapshot"]
    gl2 = next(d for d in after["deliverables"] if d["title"] == "General ledger configuration")
    assert (
        gl2["linked_tasks"] == 2
        and gl2["progress_pct"] == 50
        and gl2["hours_logged"] == 1.5
        and gl2["status"] == "in_progress"
    )
