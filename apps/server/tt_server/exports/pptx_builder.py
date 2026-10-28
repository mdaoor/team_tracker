"""PPTX export: one slide per `Slide` (plan §7.2 mapping table). Native charts for bar-shaped data, tables for the rest.

Theme colours come from the deck's token set. Speaker notes are carried over. Slide kinds with no native mapping (`image`) are
reported on the slide as a text box instead of silently dropped.
"""

from __future__ import annotations

import io
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches

from tt_server.exports import compute

from tt_server.exports.pptx_plan import PLAN_KINDS, _gantt, _plan_slide
from tt_server.exports.pptx_primitives import (
    ACCENT,
    H,
    LINE,
    MUTED,
    ROWS_PER_SLIDE,
    SEQ,
    SERIOUS,
    W,
    _bar_chart,
    _frame,
    _pages,
    _table,
    _text,
)


def build_pptx(payload: dict[str, Any], stamp_text: str) -> bytes:
    snap, sd, pres = payload["snapshot"], payload["slide_data"], payload["presentation"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    scope, as_of = snap["meta"]["scope_label"], snap["as_of"]
    slides = pres["slides"]
    # pagination makes the page count unknown until built; count after the fact via the placeholder and patch the footers
    built: list[Any] = []

    def frame(sl: dict[str, Any], **kw: Any) -> Any:
        s = _frame(prs, sl, 0, 0, stamp_text, scope, as_of, **kw)
        built.append(s)
        return s

    for sl in slides:
        kind, cfg, d = sl["kind"], sl["config"], sd.get(sl["id"]) or {}
        if kind in {"title", "section_divider"}:
            s = frame(sl, title=sl.get("title") or pres["title"])
            _text(
                s,
                Inches(0.5),
                Inches(3),
                Inches(12),
                Inches(1),
                sl.get("subtitle") or scope,
                size=20,
                color=MUTED,
            )
        elif kind == "kpis":
            s = frame(sl)
            for n, (label, val, sub) in enumerate(compute.kpis(snap)):
                x, y = Inches(0.5 + (n % 4) * 3.1), Inches(1.7 + (n // 4) * 2.3)
                box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.9), Inches(2.0))
                box.fill.background()
                box.line.color.rgb = LINE
                _text(
                    s, x + Inches(0.15), y + Inches(0.15), Inches(2.6), Inches(0.9), val, size=34, bold=True
                )
                _text(
                    s, x + Inches(0.15), y + Inches(1.0), Inches(2.6), Inches(0.4), label, size=14, bold=True
                )
                _text(
                    s, x + Inches(0.15), y + Inches(1.4), Inches(2.6), Inches(0.4), sub, size=10, color=MUTED
                )
        elif kind == "projects":
            for chunk in _pages(snap["projects"]) if snap["projects"] else [[]]:
                s = frame(sl, cont=chunk is not _pages(snap["projects"])[0])
                for n, p in enumerate(chunk):
                    its = [i for i in snap["items"] if i["project"] == p["key"]]
                    done = sum(1 for i in its if i["state"] == "done")
                    pct = round(done / len(its) * 100) if its else 0
                    x, y = Inches(0.5 + (n % 2) * 6.3), Inches(1.5 + (n // 2) * 2.7)
                    _text(
                        s,
                        x,
                        y,
                        Inches(6),
                        Inches(0.4),
                        f"{p['name']}  ({p['key']} · {p['methodology']})",
                        size=15,
                        bold=True,
                    )
                    _text(s, x, y + Inches(0.4), Inches(6), Inches(1.2), p["summary"], size=10, color=MUTED)
                    track = s.shapes.add_shape(
                        MSO_SHAPE.RECTANGLE, x, y + Inches(1.65), Inches(5.8), Inches(0.14)
                    )
                    track.fill.solid(), track.line.fill.background()
                    track.fill.fore_color.rgb = LINE
                    bar = s.shapes.add_shape(
                        MSO_SHAPE.RECTANGLE,
                        x,
                        y + Inches(1.65),
                        max(Emu(1), int(Inches(5.8) * pct / 100)),
                        Inches(0.14),
                    )
                    bar.fill.solid(), bar.line.fill.background()
                    bar.fill.fore_color.rgb = ACCENT
                    _text(
                        s,
                        x,
                        y + Inches(1.8),
                        Inches(6),
                        Inches(0.3),
                        f"{pct}% done ({done}/{len(its)})",
                        size=10,
                    )
        elif kind == "pipeline":
            s = frame(sl)
            rows = []
            exp = snap["stage_expected_days"]
            for st in snap["stages"]:
                its = [i for i in snap["items"] if i["stage"] == st]
                over = sum(
                    1
                    for i in its
                    if i["state"] != "done" and exp.get(st) and (i["stage_age_days"] or 0) > exp[st]
                )
                rows.append(
                    [
                        st,
                        str(len(its)),
                        f"{sum(i['weight'] for i in its):g}",
                        f"{over} over {exp.get(st)}d" if over else "—",
                    ]
                )
            _table(s, ["stage", "items", "Σ weight", "over expected"], rows)
        elif kind == "explorer":
            ex = compute.explorer(snap, cfg["dimA"], cfg.get("dimB"), cfg["measure"])
            s = frame(sl)
            if ex["cols"] is None:
                _bar_chart(s, ex["rows"], {cfg["measure"]: [c[0] for c in ex["cells"]]})
            else:
                mx = max([v for r in ex["cells"] for v in r] + [1])
                fills = {
                    (ri, ci + 1): SEQ[min(5, int(v / mx * 5.999))]
                    for ri, r in enumerate(ex["cells"])
                    for ci, v in enumerate(r)
                    if v
                }
                _table(
                    s,
                    [f"{cfg['dimA']} \\ {cfg['dimB']}", *ex["cols"]],
                    [
                        [r, *[str(int(v) if v else "") for v in cells]]
                        for r, cells in zip(ex["rows"], ex["cells"], strict=True)
                    ],
                    fills=fills,
                )
        elif kind == "radar":
            rd = compute.radar(snap, int(cfg["days"]))
            for k, chunk in enumerate(_pages(rd["due"])):
                s = frame(sl, cont=k > 0)
                if k == 0 and rd["overdue"]:
                    _text(
                        s,
                        Inches(0.5),
                        Inches(1.2),
                        Inches(8),
                        Inches(0.3),
                        f"{len(rd['overdue'])} overdue",
                        size=12,
                        bold=True,
                        color=SERIOUS,
                    )
                _table(
                    s,
                    ["due", "key", "item", "owner", "stage"],
                    [
                        [
                            i["planned_end"],
                            i["key"],
                            i["title"],
                            i["assignee"] or "—",
                            f"{i['stage']} {i['stage_age_days']}d",
                        ]
                        for i in chunk
                    ],
                    y=Inches(1.6),
                )
        elif kind == "gantt":
            s = frame(sl)
            _gantt(s, snap, cfg)
        elif kind == "now":
            s = frame(sl)
            _table(
                s,
                ["person", "item", "project", "since (min)", "kind"],
                [
                    [n["person"], f"{n['item']} {n['title']}", n["project"], str(n["since_min"]), n["kind"]]
                    for n in snap["now"]
                ][:ROWS_PER_SLIDE],
            )
        elif kind == "people":
            for k, chunk in enumerate(_pages(compute.people_rows(snap))):
                _table(frame(sl, cont=k > 0), compute.PEOPLE_HEADER, chunk)
        elif kind in {"violations", "habits"}:
            if d.get("restricted"):
                s = frame(sl)
                _text(
                    s,
                    Inches(0.5),
                    Inches(3),
                    Inches(12),
                    Inches(1),
                    "Restricted: you do not have access to this data.",
                    size=20,
                    color=MUTED,
                )
                continue
            rows = (
                [
                    [
                        v["severity"],
                        v["target"],
                        v["policy"].replace("_", " "),
                        v["detail"],
                        v["person"] or "",
                    ]
                    for v in compute.scoped_violations(snap, snap["items"])
                ]
                if kind == "violations"
                else [
                    [
                        h["subject"],
                        h["key"].replace("_", " "),
                        h["polarity"],
                        f"{round(h['rate'] * 100)}%",
                        h["trend"],
                    ]
                    for h in snap["habits"]
                ]
            )
            hdr = (
                ["severity", "target", "policy", "detail", "person"]
                if kind == "violations"
                else ["subject", "habit", "polarity", "rate", "trend"]
            )
            for k, chunk in enumerate(_pages(rows)):
                s = frame(sl, cont=k > 0)
                if not rows:
                    _text(
                        s,
                        Inches(0.5),
                        Inches(2),
                        Inches(12),
                        Inches(1),
                        "No open violations in scope." if kind == "violations" else "No habits in scope.",
                        size=18,
                        color=MUTED,
                    )
                else:
                    _table(s, hdr, chunk)
        elif kind == "chart":
            s = frame(sl)
            if d.get("cols"):
                mx = max([v for r in d["cells"] for v in r] + [1])
                fills = {
                    (ri, ci + 1): SEQ[min(5, int(v / mx * 5.999))]
                    for ri, r in enumerate(d["cells"])
                    for ci, v in enumerate(r)
                    if v
                }
                _table(
                    s,
                    [cfg["rows"], *d["cols"]],
                    [
                        [r["value"], *[f"{v:g}" if v else "" for v in c]]
                        for r, c in zip(d["rows"], d["cells"], strict=True)
                    ],
                    fills=fills,
                )
            else:
                _bar_chart(
                    s,
                    [r["value"] for r in d.get("rows", [])],
                    {cfg.get("measure", "count"): [c[0] for c in d.get("cells", [])]},
                )
        elif kind in {"table", "view"}:
            items = d.get("items", [])
            for k, chunk in enumerate(_pages(items)):
                _table(
                    frame(sl, title=sl.get("title") or (d.get("view") or {}).get("name"), cont=k > 0),
                    ["key", "title", "stage", "priority", "weight", "assignee"],
                    [
                        [
                            i["key"],
                            i["title"],
                            i["stage"] or "",
                            i["priority"],
                            f"{i['weight']:g}",
                            i["assignee"] or "",
                        ]
                        for i in chunk
                    ],
                )
        elif kind in {"insight", "project_detail"}:
            s = frame(sl)
            ins = d.get("insights", [])
            _text(
                s,
                Inches(0.5),
                Inches(1.5),
                Inches(12),
                Inches(5),
                "\n\n".join(f"{i['project']}: {i['body']}" for i in ins)
                or "\n\n".join(
                    f"{p['key']}: {p['summary']}"
                    for p in snap["projects"]
                    if kind == "insight" or p["key"] == cfg.get("project_key")
                ),
                size=14,
            )
        elif kind in PLAN_KINDS:
            _plan_slide(frame, sl, kind, cfg, snap)
        elif kind == "markdown":
            _text(frame(sl), Inches(0.5), Inches(1.5), Inches(12), Inches(5), cfg["text"], size=16)
        elif kind == "image":
            _text(
                frame(sl),
                Inches(0.5),
                Inches(3),
                Inches(12),
                Inches(1),
                f"[image: {cfg.get('url') or cfg.get('attachment_id') or 'none'}]",
                size=16,
                color=MUTED,
            )
    total = len(built)
    for n, s in enumerate(built, 1):  # footers carry the page counter now that pagination is known
        for shape in s.shapes:
            if shape.has_text_frame and shape.text_frame.text.endswith(" · 0/0"):
                shape.text_frame.paragraphs[0].runs[0].text = (
                    shape.text_frame.text[: -len("0/0")] + f"{n}/{total}"
                )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
