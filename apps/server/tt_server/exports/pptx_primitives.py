"""Shared PPTX drawing primitives: theme constants, text boxes, the slide frame, tables, native bar charts and pagination."""

from __future__ import annotations

from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

W, H = Inches(13.333), Inches(7.5)
INK, MUTED, LINE = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0xE3, 0xE1, 0xDC)
SEQ = [RGBColor.from_string(c) for c in ("CDE2FB", "9EC5F4", "6DA7EC", "3987E5", "1C5CAB", "0D366B")]
ACCENT, SERIOUS = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEC, 0x83, 0x5A)
ROWS_PER_SLIDE = 12


def _text(
    slide: Any,
    x: Any,
    y: Any,
    w: Any,
    h: Any,
    text: str,
    *,
    size: int = 14,
    bold: bool = False,
    color: RGBColor = INK,
) -> Any:
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return tb


def _frame(
    prs: Any,
    sl: dict[str, Any],
    page: int,
    total: int,
    stamp_text: str,
    scope: str,
    as_of: str,
    title: str | None = None,
    cont: bool = False,
) -> Any:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = (title or sl.get("title") or sl["kind"].replace("_", " ").title()) + (" (cont.)" if cont else "")
    _text(s, Inches(0.5), Inches(0.3), Inches(10), Inches(0.7), t, size=26, bold=True)
    _text(
        s,
        Inches(10.3),
        Inches(0.4),
        Inches(2.6),
        Inches(0.4),
        f"{scope} · as of {as_of[:16].replace('T', ' ')}",
        size=10,
        color=MUTED,
    )
    if sl.get("subtitle"):
        _text(s, Inches(0.5), Inches(0.95), Inches(12), Inches(0.4), sl["subtitle"], size=14, color=MUTED)
    _text(
        s,
        Inches(0.5),
        Inches(7.05),
        Inches(11.5),
        Inches(0.3),
        f"{stamp_text} · {page}/{total}",
        size=9,
        color=MUTED,
    )
    if sl.get("notes"):
        s.notes_slide.notes_text_frame.text = sl["notes"]
    return s


def _table(
    slide: Any,
    header: list[str],
    rows: list[list[str]],
    *,
    y: Any = None,
    fills: dict[tuple[int, int], RGBColor] | None = None,
) -> None:
    y = y if y is not None else Inches(1.5)
    shape = slide.shapes.add_table(
        len(rows) + 1, len(header), Inches(0.5), y, Inches(12.3), Inches(0.35) * (len(rows) + 1)
    )
    tbl = shape.table
    for c, h in enumerate(header):
        tbl.cell(0, c).text = str(h)
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = str(v)
            if fills and (r - 1, c) in fills:
                tbl.cell(r, c).fill.solid()
                tbl.cell(r, c).fill.fore_color.rgb = fills[(r - 1, c)]
    for r in range(len(rows) + 1):
        for c in range(len(header)):
            for p in tbl.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11)


def _bar_chart(slide: Any, cats: list[str], series: dict[str, list[float]], *, y: Any = None) -> None:
    y = y if y is not None else Inches(1.5)
    if not cats or not series or not any(series.values()):  # python-pptx refuses an empty category axis
        box = slide.shapes.add_textbox(Inches(0.5), y, Inches(12.3), Inches(1))
        box.text_frame.text = "No data in scope."
        return
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series.items():
        cd.add_series(name, vals)
    slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), y, Inches(12.3), Inches(5.3), cd)


def _pages(rows: list[Any]) -> list[list[Any]]:
    return [rows[i : i + ROWS_PER_SLIDE] for i in range(0, len(rows), ROWS_PER_SLIDE)] or [[]]
