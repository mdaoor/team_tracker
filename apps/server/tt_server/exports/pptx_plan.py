"""PPTX renderers for the plan slide kinds (roadmap, milestones, deliverables, ...) and the epic Gantt."""

from __future__ import annotations

from datetime import date
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches


from tt_server.exports.pptx_primitives import ACCENT, LINE, MUTED, SERIOUS, _pages, _table, _text


PLAN_KINDS = {"roadmap", "milestones", "deliverables", "plan_vs_actual", "phase_status", "project_overview"}
STATUS_FILL = {
    "overdue": RGBColor(0xEC, 0x83, 0x5A),
    "at_risk": RGBColor(0xF2, 0xB7, 0x5A),
    "done": RGBColor(0x3C, 0xA6, 0x6B),
}


def _plan_slide(frame: Any, sl: dict[str, Any], kind: str, cfg: dict[str, Any], snap: dict[str, Any]) -> None:
    """Plan kinds: tables from the manager-entered plan; the roadmap as shapes on a date axis. Empty plans say so."""
    pk = cfg.get("project_key")
    phases = [p for p in snap.get("phases", []) if not pk or p["project"] == pk]
    delivs = [
        d
        for d in snap.get("deliverables", [])
        if (not pk or d["project"] == pk) and (not cfg.get("phase") or d["phase"] == cfg["phase"])
    ]
    mss = [x for x in snap.get("milestones", []) if not pk or x["project"] == pk]
    projects = [
        p
        for p in snap["projects"]
        if (pk and p["key"] == pk) or (not pk and any(x["project"] == p["key"] for x in phases + mss))
    ]
    empty = {
        "roadmap": phases,
        "milestones": mss,
        "deliverables": delivs,
        "plan_vs_actual": delivs,
        "phase_status": phases,
        "project_overview": projects,
    }[kind]
    if not empty:
        _text(
            frame(sl),
            Inches(0.5),
            Inches(1.6),
            Inches(12),
            Inches(1),
            "No plan for this scope yet.",
            size=16,
            color=MUTED,
        )
        return
    dash = lambda v: "—" if v in (None, "") else str(v)  # noqa: E731
    if kind == "project_overview":
        s = frame(sl)
        y: int = int(Inches(1.5))
        for p in projects:
            nxt = next((x for x in mss if x["project"] == p["key"] and x["status"] != "completed"), None)
            _text(
                s,
                Inches(0.5),
                y,
                Inches(12),
                Inches(0.5),
                f"{p['name']} ({p['key']})  ·  health: {dash(p.get('health'))}",
                size=18,
                bold=True,
            )
            _text(
                s, Inches(0.5), y + Inches(0.5), Inches(12), Inches(1.4),
                f"Owner {dash(p.get('owner'))}\nDates {dash(p.get('planned_start'))} → {dash(p.get('planned_end'))}\n"
                f"Status {dash(p.get('status'))}   Baseline {dash(p.get('baseline_name'))}\n"
                f"Next milestone {nxt['name'] + ' (' + nxt['due'] + ')' if nxt else '—'}\n{p.get('summary') or ''}",
                size=12,
            )  # fmt: skip
            y += int(Inches(2.3))
    elif kind == "milestones":
        _table(
            frame(sl),
            ["Milestone", "Kind", "Due", "Status", "Deliverables"],
            [
                [
                    x["name"],
                    x["kind"],
                    x["due"],
                    x["status"],
                    f"{x['deliverables_done']}/{x['deliverables_total']}" if x["deliverables_total"] else "—",
                ]
                for x in mss
            ],
        )
    elif kind == "phase_status":
        _table(
            frame(sl),
            ["Phase", "Dates", "Status", "Progress", "Deliverables"],
            [
                [
                    p["title"],
                    f"{dash(p['planned_start'])} → {dash(p['planned_end'])}",
                    p["status"],
                    f"{p['progress_pct']}%" if p["progress_pct"] is not None else "—",
                    p["deliverable_count"],
                ]
                for p in phases
            ],
        )
    elif kind in {"deliverables", "plan_vs_actual"}:
        for k, chunk in enumerate(_pages(delivs)):
            s = frame(sl, cont=k > 0)
            if kind == "deliverables":
                _table(
                    s,
                    ["Deliverable", "Owner", "Planned", "Status", "Progress", "Tasks", "Hours"],
                    [
                        [
                            d["title"],
                            dash(d["owner"]),
                            f"{dash(d['planned_start'])} → {dash(d['planned_end'])}",
                            d["status"],
                            f"{d['progress_pct']}%" if d["progress_pct"] is not None else "—",
                            d["linked_tasks"] or "—",
                            d["hours_logged"] or "—",
                        ]
                        for d in chunk
                    ],
                )
            else:
                _table(
                    s,
                    ["Deliverable", "Baseline end", "Current end", "Variance (d)", "Actual", "Status"],
                    [
                        [
                            d["title"],
                            dash(d["baseline_end"]),
                            dash(d["planned_end"]),
                            dash(d["baseline_variance_days"]),
                            dash(d["actual_end"]),
                            d["status"],
                        ]
                        for d in chunk
                    ],
                )
    else:  # roadmap: bars on a date axis (native shapes, so it stays editable)
        s = frame(sl)
        rows = [("phase", p) for p in phases for p in [p]]
        rows = [
            r
            for p in phases
            for r in [("phase", p), *[("deliverable", d) for d in delivs if d["phase"] == p["key"]]]
        ]
        stamps = [
            date.fromisoformat(v[:10]) for _, r in rows for v in (r["planned_start"], r["planned_end"]) if v
        ] + [date.fromisoformat(x["due"][:10]) for x in mss]
        if not stamps:
            return
        lo, hi = min(stamps), max(stamps)
        span = max((hi - lo).days, 7)
        left, width, top = Inches(3.4), Inches(9.4), Inches(1.5)
        step = min(int(Inches(0.32)), int(Inches(5.2) / max(len(rows), 1)))
        for i, (rk, r) in enumerate(rows[:18]):
            y = top + step * i
            _text(
                s,
                Inches(0.5),
                y,
                Inches(2.9),
                step,
                ("" if rk == "phase" else "   ") + r["title"],
                size=10,
                bold=rk == "phase",
            )
            if r["planned_start"] and r["planned_end"]:
                a, b = date.fromisoformat(r["planned_start"][:10]), date.fromisoformat(r["planned_end"][:10])
                x = left + int(width * (a - lo).days / span)
                w = max(Emu(Inches(0.05)), int(width * (b - a).days / span))
                bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + int(step * 0.15), w, int(step * 0.6))
                bar.fill.solid(), bar.line.fill.background()
                bar.fill.fore_color.rgb = STATUS_FILL.get(r["status"], ACCENT)
        for x in mss:
            xx = left + int(width * (date.fromisoformat(x["due"][:10]) - lo).days / span)
            m_ = s.shapes.add_shape(
                MSO_SHAPE.DIAMOND, xx - Inches(0.08), Inches(1.2), Inches(0.16), Inches(0.16)
            )
            m_.fill.solid(), m_.line.fill.background()
            m_.fill.fore_color.rgb = SERIOUS if x["status"] == "overdue" else RGBColor(0x8F, 0x5B, 0xD6)
        _text(
            s,
            left,
            Inches(6.9),
            width,
            Inches(0.3),
            f"{lo.isoformat()}  →  {hi.isoformat()}   ◆ milestones",
            size=9,
            color=MUTED,
        )


def _gantt(slide: Any, snap: dict[str, Any], cfg: dict[str, Any]) -> None:
    rows = []
    for e in snap["epics"]:
        its = [i for i in snap["items"] if i["epic"] == e["key"] and i["planned_start"] and i["planned_end"]]
        if its:
            rows.append(
                (
                    f"{e['key']} {e['title']}"[:34],
                    min(i["planned_start"] for i in its),
                    max(i["planned_end"] for i in its),
                    sum(1 for i in its if i["state"] == "done"),
                    len(its),
                )
            )
    if not rows:
        _text(
            slide,
            Inches(0.5),
            Inches(2),
            Inches(12),
            Inches(1),
            "No scheduled epics in scope.",
            size=18,
            color=MUTED,
        )
        return
    d0 = min(date.fromisoformat(r[1]) for r in rows)
    d1 = max(date.fromisoformat(r[2]) for r in rows)
    span = max(1, (d1 - d0).days)
    x0, wd = Inches(3.6), Inches(9.2)
    for n, (label, a, b, done, tot) in enumerate(rows[:14]):
        y = Inches(1.6 + n * 0.38)
        _text(slide, Inches(0.5), y - Inches(0.05), Inches(3), Inches(0.3), label, size=10)
        xa = x0 + int(wd * (date.fromisoformat(a) - d0).days / span)
        wbar = max(Emu(Inches(0.05)), int(wd * ((date.fromisoformat(b) - date.fromisoformat(a)).days) / span))
        track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, xa, y, wbar, Inches(0.22))
        track.fill.solid(), track.line.fill.background()
        track.fill.fore_color.rgb = LINE
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, xa, y, max(Emu(1), int(wbar * done / max(1, tot))), Inches(0.22)
        )
        bar.fill.solid(), bar.line.fill.background()
        bar.fill.fore_color.rgb = ACCENT
    today = date.fromisoformat(snap["as_of"][:10])
    if d0 <= today <= d1:
        ln = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            x0 + int(wd * (today - d0).days / span),
            Inches(1.5),
            Emu(12000),
            Inches(0.38 * min(len(rows), 14) + 0.2),
        )
        ln.fill.solid(), ln.line.fill.background()
        ln.fill.fore_color.rgb = SERIOUS
